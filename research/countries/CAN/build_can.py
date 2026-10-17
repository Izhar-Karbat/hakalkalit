"""Canada: average price of a home sold, annual, 2000-2025, Canadian dollars.

Source of every number: files in raw/, downloaded from the Canadian Real Estate Association (CREA).
    main series   raw/crea_forecast_data_Oct_2026_annual_avg_price_1980-2025.xlsx, sheet "Chart",
                  column "Residential Average Sale Price" (annual, MLS® Systems, whole country)
    cross-check 1 the same file's 2026 forecast row must equal the dollar figure CREA prints in the
                  October 2026 forecast release (raw/crea_quarterly_forecast_page_Oct_2026.html),
                  which proves the file is the one behind CREA's published text
    cross-check 2 raw/crea_MLS_HPI_Oct_2026.zip, "Not Seasonally Adjusted (A).xlsx", sheet AGGREGATE,
                  Composite_Benchmark: CREA's quality-adjusted benchmark price, 2005 onward

Run: python3 build_can.py  (needs openpyxl). Reads only raw/, writes can_avg_dwelling_price_2000-2025.csv.
"""
import csv
import html
import io
import re
import zipfile
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
RAW = HERE / "raw"
OUT = HERE / "can_avg_dwelling_price_2000-2025.csv"

FORECAST_XLSX = RAW / "crea_forecast_data_Oct_2026_annual_avg_price_1980-2025.xlsx"
FORECAST_PAGE = RAW / "crea_quarterly_forecast_page_Oct_2026.html"
HPI_ZIP = RAW / "crea_MLS_HPI_Oct_2026.zip"

YEARS = range(2000, 2026)
NOTE = ("CREA MLS(R) residential average sale price, annual, Canada; mean of all residential "
        "sales on MLS(R) Systems (mostly resale); not quality-adjusted; industry body, not a "
        "government statistic")


def read_avg_price():
    wb = openpyxl.load_workbook(FORECAST_XLSX, data_only=True, read_only=True)
    rows = list(wb["Chart"].iter_rows(values_only=True))
    header = rows[0]
    assert header[2] == "Residential Average Sale Price", header
    actual, forecast = {}, {}
    for label, _sales, price in rows[1:]:
        if label is None or price is None:
            continue
        label = str(label).strip()
        if label.endswith("F"):
            forecast[int(label[:-1])] = float(price)
        else:
            actual[int(label)] = float(price)
    return actual, forecast


def check_forecast_against_release(forecast):
    """The release prints the 2026 forecast to the dollar; the file must agree."""
    text = html.unescape(re.sub(r"<[^>]+>", " ", FORECAST_PAGE.read_text(errors="ignore")))
    text = re.sub(r"\s+", " ", text)
    m = re.search(r"national average home price is forecast to (?:[^.]|\.\d)*? to \$([\d,]+) in 2026", text)
    assert m, "2026 forecast sentence not found in release page"
    printed = int(m.group(1).replace(",", ""))
    assert round(forecast[2026]) == printed, (forecast[2026], printed)
    pct = re.search(r"national average home price is forecast to (?:[^.]|\.\d)*?by just ([\d.]+)% on an annual basis to \$[\d,]+ in 2026", text)
    return printed, (float(pct.group(1)) if pct else None)


def read_hpi_benchmark():
    with zipfile.ZipFile(HPI_ZIP) as z:
        data = z.read("Not Seasonally Adjusted (A).xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    rows = list(wb["AGGREGATE"].iter_rows(values_only=True))
    col = rows[0].index("Composite_Benchmark")
    return {int(r[0]): float(r[col]) for r in rows[1:] if r[0] is not None and r[col] is not None}


def main():
    actual, forecast = read_avg_price()
    printed_2026, pct = check_forecast_against_release(forecast)
    hpi = read_hpi_benchmark()

    with OUT.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["year", "avg_dwelling_price_local", "currency", "note"])
        for y in YEARS:
            w.writerow([y, round(actual[y]) if y in actual else "", "CAD", NOTE if y in actual else "not sourced"])

    print(f"wrote {OUT.name}: {sum(y in actual for y in YEARS)} of {len(YEARS)} years")
    print(f"check 1: file 2026F = {forecast[2026]:,.0f}; release prints ${printed_2026:,} -> match")
    if pct is not None:
        implied = printed_2026 / (1 + pct / 100)
        lo, hi = printed_2026 / (1 + (pct + 0.05) / 100), printed_2026 / (1 + (pct - 0.05) / 100)
        print(f"         release says +{pct}% on 2025 -> 2025 implied {implied:,.0f} "
              f"(rounding band {lo:,.0f}-{hi:,.0f}); file 2025 = {actual[2025]:,.0f} "
              f"-> {'inside' if lo <= actual[2025] <= hi else 'OUTSIDE'} band")
    print("check 1b (revisions): base year implied by each January forecast release vs today's file")
    for fname, base in (("crea_forecast_release_2025-01-15.html", 2024),
                        ("crea_forecast_release_2026-01-15.html", 2025)):
        t = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", (RAW / fname).read_text(errors="ignore"))))
        m = re.search(r"national average home price is forecast to (?:[^.]|\.\d)*?by ([\d.]+)% on an "
                      r"annual basis to \$([\d,]+) in " + str(base + 1), t)
        p, v = float(m.group(1)), int(m.group(2).replace(",", ""))
        lo, hi = v / (1 + (p + 0.05) / 100), v / (1 + (p - 0.05) / 100)
        print(f"   {fname}: {base + 1}F ${v:,} at +{p}% -> {base} then {lo:,.0f}-{hi:,.0f}; "
              f"file now {actual[base]:,.0f} ({(actual[base] / (v / (1 + p / 100)) - 1) * 100:+.2f}%)")
    print("check 2: average price vs HPI composite benchmark (Aggregate)")
    for y in (2005, 2007, 2016, 2021, 2025):
        print(f"   {y}: average {actual[y]:>9,.0f}  benchmark {hpi[y]:>9,.0f}  "
              f"ratio {actual[y] / hpi[y]:.3f}")
    print(f"   growth 2005->2025: average x{actual[2025] / actual[2005]:.2f}, "
          f"benchmark x{hpi[2025] / hpi[2005]:.2f}")
    for y in (2000, 2007, 2016, 2025):
        print(f"   {y}: {actual[y]:,.0f} CAD")


if __name__ == "__main__":
    main()
