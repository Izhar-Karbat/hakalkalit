"""Rebuild the apartment-price-in-years-of-salary table from the CBS public API,
and check it against the table we published.

Pipeline
--------
1. FETCH   the CBS API URLs below (the same links sent to Globes, plus a
           1994–1999 extension) and save each raw JSON response to raw/.
2. PARSE   the JSON into {year: value} series, asserting everything the
           calculation relies on (single index base, complete months, the
           expected frequency and unit per wage series, no truncated paging).
3. BUILD   the table:
             price  = annual mean of index 40010 × (₪ per index point)
             wage   = CBS average monthly wage per employee post, NII-based
             ratio  = price ÷ (wage × 12)
4. COMPARE every cell with israel_housing_CBS_workings_1994-2025.csv and
           "apartment prices 2000.xlsx" when those files sit next to this script.
           Exit status 1 on any mismatch. Without them (e.g. a copy taken from the
           web page) the table is built and printed, and the comparison is skipped.

Run:  python3 rebuild_from_cbs.py            (fetches live, overwrites raw/)
      python3 rebuild_from_cbs.py --offline  (re-parses the saved raw/ files)
"""
import csv
import json
import statistics
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
RAW = HERE / "raw"

# --- Sources -----------------------------------------------------------------
# The two links sent to Globes cover 2000–2025.
PRICE_URL = ("https://api.cbs.gov.il/index/data/price?id=40010&format=json&download=false"
             "&startPeriod=01-2000&endPeriod=12-2025&PageSize=400")
WAGE_URL = ("https://apis.cbs.gov.il/series/data/list?id=30014,45002,45637,615908"
            "&startperiod=01-2000&endperiod=12-2025&format=json&download=false")
# 1994–1999 extension: same index; wage series 30014 starts in 1995, so 1994 is
# derived from wage index 30245 (1994 = 100).
PRICE_URL_EARLY = ("https://api.cbs.gov.il/index/data/price?id=40010&format=json&download=false"
                   "&startPeriod=01-1994&endPeriod=12-1999&PageSize=400")
WAGE_URL_EARLY = ("https://apis.cbs.gov.il/series/data/list?id=30014,30245"
                  "&startperiod=01-1994&endperiod=12-1999&format=json&download=false&PageSize=200")
# The anchor window's index months come from PRICE_URL, so it must reach Mar 2025.

# --- Price anchor --------------------------------------------------------------
# The index gives movement, not a shekel level. The level is anchored to CBS
# average national prices (₪ thousands; free-market deals, 1–6 rooms, not
# quality-adjusted), matched to the index average over the same 12 months.
# CBS revises each quarter once, in the next quarter's release; each value below
# is the revised figure, as printed in section 4 of (Q2:2024 as the year-ago
# comparison, the others as "the previous quarter"):
#   Q2:2024, Q1:2025  release 257/2025 (15.8.2025)
#     https://www.cbs.gov.il/he/mediarelease/Madad/DocLib/2025/257/10_25_257b.pdf
#   Q3:2024           release 051/2025 (14.2.2025)  [first print 2,242.9 in 358/2024]
#     https://www.cbs.gov.il/he/mediarelease/Madad/DocLib/2025/051/10_25_051b.pdf
#   Q4:2024           release 151/2025 (15.5.2025)  [first print 2,330.9 in 051/2025]
#     https://www.cbs.gov.il/he/mediarelease/Madad/DocLib/2025/151/10_25_151b.pdf
# Local copies: cbs_release_*_avg_prices.pdf in this folder.
ANCHOR_QUARTERS_K = {"Q2:2024": 2330.0, "Q3:2024": 2270.4, "Q4:2024": 2350.9, "Q1:2025": 2350.9}
ANCHOR_MONTHS = [(2024, m) for m in range(4, 13)] + [(2025, m) for m in range(1, 4)]

# --- Wage series, one per period (CBS re-issues the series per processing system) --
#   id      years       frequency  annual value used
WAGE_PLAN = [
    ("30014", 1995, 2001, "חודש", "mean of 12 months"),   # 1994 processing system
    ("45002", 2002, 2003, "חודש", "mean of 12 months"),   # 2001 processing system
    ("45637", 2004, 2010, "שנה", "published annual"),     # 2005 processing system
    ("615908", 2011, 2025, "שנה", "published annual"),    # 2011 industry classification
]
WAGE_INDEX_1994 = "30245"   # wage index, 1994 = 100; used only to derive 1994


# =============================== 1. FETCH =====================================
def fetch(url, name, offline):
    """Return the parsed JSON for url, saving the raw response to raw/<name>.json."""
    path = RAW / f"{name}.json"
    if not offline:
        RAW.mkdir(exist_ok=True)
        body = subprocess.run(["curl", "-s", "-m", "90", url],
                              capture_output=True, text=True, check=True).stdout
        path.write_text(body, encoding="utf-8")
    return json.loads(path.read_text(encoding="utf-8"))


# =============================== 2. PARSE =====================================
def parse_price(doc, y0, y1):
    """Index 40010 JSON -> {(year, month): value}. Asserts one page, one base, all months."""
    assert doc["paging"]["last_page"] == 1, "price response is paginated: data truncated"
    series = doc["month"][0]
    assert series["code"] == 40010
    bases = {d["currBase"]["baseDesc"] for d in series["date"]}
    assert bases == {"1993 ממוצע"}, f"unexpected index base(s): {bases}"
    out = {(d["year"], d["month"]): d["currBase"]["value"] for d in series["date"]}
    missing = [(y, m) for y in range(y0, y1 + 1) for m in range(1, 13) if (y, m) not in out]
    assert not missing, f"missing index months: {missing}"
    return out


def parse_wages(doc):
    """Series-list JSON -> {series_id: (frequency_name, unit_name, {period 'YYYY-MM': value})}."""
    assert doc["DataSet"]["paging"]["last_page"] == 1, "wage response is paginated: data truncated"
    out = {}
    for s in doc["DataSet"]["Series"]:
        out[str(s["id"])] = (s["time"]["name"], s["unit"]["name"],
                             {o["TimePeriod"]: o["Value"] for o in s["obs"]})
    return out


def annual_wage(wages, sid, year, freq):
    """Annual average monthly wage for one year from one series."""
    got_freq, unit, obs = wages[sid]
    assert got_freq == freq, f"{sid}: expected frequency {freq}, got {got_freq}"
    assert unit == 'ש"ח', f"{sid}: expected unit ש\"ח, got {unit}"
    if freq == "חודש":
        months = [obs[f"{year}-{m:02d}"] for m in range(1, 13)]   # KeyError if a month is missing
        return statistics.mean(months)
    return obs[f"{year}-01"]                                       # annual series are dated YYYY-01


# =============================== 3. BUILD =====================================
def build(offline):
    price = parse_price(fetch(PRICE_URL, "price_40010_2000-2025", offline), 2000, 2025)
    price |= parse_price(fetch(PRICE_URL_EARLY, "price_40010_1994-1999", offline), 1994, 1999)
    wages = parse_wages(fetch(WAGE_URL, "wages_2000-2025", offline))
    early = parse_wages(fetch(WAGE_URL_EARLY, "wages_1994-1999", offline))
    for sid, (f, u, obs) in early.items():               # merge 1994–1999 observations
        if sid in wages:
            wages[sid][2].update(obs)
        else:
            wages[sid] = (f, u, obs)

    anchor_price = statistics.mean(ANCHOR_QUARTERS_K.values()) * 1000
    anchor_index = statistics.mean(price[k] for k in ANCHOR_MONTHS)
    per_point = anchor_price / anchor_index

    rows = []
    for year in range(1994, 2026):
        index_avg = statistics.mean(price[(year, m)] for m in range(1, 13))
        if year == 1994:
            growth_95 = statistics.mean(wages[WAGE_INDEX_1994][2][f"1995-{m:02d}"]
                                        for m in range(1, 13)) / 100
            wage = annual_wage(wages, "30014", 1995, "חודש") / growth_95
            source = "derived: 30014 (1995) / 30245"
        else:
            sid, _, _, freq, _ = next(p for p in WAGE_PLAN if p[1] <= year <= p[2])
            wage = annual_wage(wages, sid, year, freq)
            source = sid
        p = round(index_avg * per_point, -2)
        rows.append({"year": year, "index": index_avg, "price": int(p), "wage": wage,
                     "ratio": round(p / (wage * 12), 1), "source": source})
    return rows, anchor_price, anchor_index, per_point


# =============================== 4. COMPARE ===================================
def load_published():
    workings = HERE / "israel_housing_CBS_workings_1994-2025.csv"
    if not workings.exists():
        print("israel_housing_CBS_workings_1994-2025.csv not found: building the table without the comparison")
        return {}, {}
    csv_rows = {int(r["year"]): r for r in csv.DictReader(open(workings, encoding="utf-8-sig"))}
    xlsx = {}
    try:
        import openpyxl
        if not (HERE / "apartment prices 2000.xlsx").exists():
            return csv_rows, xlsx
        ws = openpyxl.load_workbook(HERE / "apartment prices 2000.xlsx").active
        for r in range(2, ws.max_row + 1):
            if isinstance(ws[f"A{r}"].value, int):
                xlsx[ws[f"A{r}"].value] = (ws[f"B{r}"].value, ws[f"C{r}"].value, ws[f"D{r}"].value)
    except ImportError:
        print("openpyxl not installed: skipping the xlsx comparison")
    return csv_rows, xlsx


def main():
    offline = "--offline" in sys.argv
    rows, anchor_price, anchor_index, per_point = build(offline)
    csv_rows, xlsx = load_published()

    print(f"Anchor: mean of {list(ANCHOR_QUARTERS_K)} = ₪{anchor_price:,.0f}; "
          f"index mean Apr-2024..Mar-2025 = {anchor_index:.3f}; ₪ per point = {per_point:.4f}\n")
    print(f"{'year':<5}{'index':>9}{'price':>11}{'wage':>9}{'ratio':>7}  {'series':<8} csv   xlsx")
    mismatches = 0
    for r in rows:
        c = csv_rows.get(r["year"])
        ok_csv = None if c is None else (abs(r["index"] - float(c["CBS_40010_HPI_annual_avg_1993=100"])) < 0.0051
                  and r["price"] == int(c[next(k for k in c if k.startswith("avg_price"))])
                  and round(r["wage"]) == int(c["avg_monthly_wage_ILS"])
                  and r["ratio"] == float(c["ratio_years"]))
        x = xlsx.get(r["year"])
        ok_x = None if x is None else (r["price"] == x[0] and round(r["wage"]) == x[1]
                                       and r["ratio"] == float(x[2]))
        mismatches += (ok_csv is False) + (ok_x is False)
        mark = lambda ok: "—" if ok is None else ("ok" if ok else "DIFF")
        print(f"{r['year']:<5}{r['index']:>9.3f}{r['price']:>11,}{r['wage']:>9.1f}{r['ratio']:>7.1f}"
              f"  {r['source'][:8]:<8} {mark(ok_csv):<5} {mark(ok_x)}")
    print(f"\nMismatches: {mismatches}")
    sys.exit(1 if mismatches else 0)


if __name__ == "__main__":
    main()
